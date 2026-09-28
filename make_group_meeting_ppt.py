#!/usr/bin/env python3
"""Generate a group-meeting PPTX from the Fudan template."""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
WORKSPACE = ROOT.parent
TEMPLATE = Path("/Users/gjz/Downloads/复旦ppt模板/复旦大学ppt模板2.pptx")
OUT_DIR = ROOT / "outputs"
IMG1 = Path("/Users/gjz/Pictures/检测1.jpg")
IMG2 = Path("/Users/gjz/Pictures/检测2.png")

TITLE = "复杂环境施工安全快速检测系统"
SUBTITLE = "多模型并行投票 · 真实法规条款 · 8秒内快速初筛"


FUDAN_RED = RGBColor(130, 0, 35)
FUDAN_DARK = RGBColor(35, 35, 35)
FUDAN_GOLD = RGBColor(188, 150, 66)
FUDAN_GRAY = RGBColor(95, 95, 95)
LIGHT_BG = RGBColor(247, 247, 245)
PANEL_BG = RGBColor(255, 255, 255)
PALE_RED = RGBColor(246, 232, 236)
PALE_GOLD = RGBColor(251, 245, 226)
GREEN = RGBColor(38, 132, 92)
ORANGE = RGBColor(207, 117, 32)
BLUE = RGBColor(47, 97, 155)


def delete_all_slides(prs: Presentation) -> None:
    slide_id_list = prs.slides._sldIdLst
    for slide_id in list(slide_id_list):
        r_id = slide_id.rId
        prs.part.drop_rel(r_id)
        slide_id_list.remove(slide_id)


def set_text(
    shape,
    text: str,
    size: int = 22,
    color: RGBColor = FUDAN_DARK,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    font: str = "PingFang SC",
    line_spacing: float | None = None,
) -> None:
    shape.text_frame.clear()
    shape.text_frame.word_wrap = True
    shape.text_frame.margin_left = Inches(0.04)
    shape.text_frame.margin_right = Inches(0.04)
    shape.text_frame.margin_top = Inches(0.02)
    shape.text_frame.margin_bottom = Inches(0.02)
    p = shape.text_frame.paragraphs[0]
    p.alignment = align
    if line_spacing:
        p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_textbox(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    size: int = 22,
    color: RGBColor = FUDAN_DARK,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    font: str = "PingFang SC",
    line_spacing: float | None = None,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    set_text(shape, text, size=size, color=color, bold=bold, align=align, font=font, line_spacing=line_spacing)
    return shape


def add_rect(slide, x, y, w, h, fill=PANEL_BG, line=None, radius=False):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def add_badge(slide, x, y, text, fill=PALE_RED, color=FUDAN_RED, w=None):
    width = w if w is not None else max(1.2, 0.23 * len(text))
    shape = add_rect(slide, x, y, width, 0.36, fill=fill, line=None, radius=True)
    set_text(shape, text, size=12, color=color, bold=True, align=PP_ALIGN.CENTER)
    return shape


def add_footer(slide, idx: int, total: int) -> None:
    add_rect(slide, 0, 7.18, 13.33, 0.04, fill=FUDAN_RED)
    add_textbox(slide, 0.55, 7.22, 6.5, 0.18, "复杂环境安全巡检 · 组会汇报", size=8, color=FUDAN_GRAY)
    add_textbox(slide, 11.75, 7.22, 1.0, 0.18, f"{idx:02d}/{total:02d}", size=8, color=FUDAN_GRAY, align=PP_ALIGN.RIGHT)


def title_slide(slide, kicker: str, title: str, subtitle: str, date_text: str) -> None:
    add_rect(slide, 0, 0, 13.33, 7.5, fill=LIGHT_BG)
    add_rect(slide, 0, 0, 0.28, 7.5, fill=FUDAN_RED)
    add_rect(slide, 0.28, 0, 0.07, 7.5, fill=FUDAN_GOLD)
    add_badge(slide, 0.78, 0.76, kicker, fill=PALE_RED, color=FUDAN_RED, w=2.2)
    add_textbox(slide, 0.78, 1.45, 9.2, 0.75, title, size=36, color=FUDAN_RED, bold=True)
    add_textbox(slide, 0.82, 2.35, 8.6, 0.6, subtitle, size=20, color=FUDAN_DARK)
    add_textbox(slide, 0.82, 5.9, 8.5, 0.45, "面向施工现场单图安全风险快速初筛，输出风险点、真实可查条例与实名投票证据。", size=16, color=FUDAN_GRAY)
    add_textbox(slide, 0.82, 6.54, 5.0, 0.25, date_text, size=11, color=FUDAN_GRAY)
    add_rect(slide, 9.75, 0.75, 2.75, 5.55, fill=FUDAN_RED, radius=True)
    add_textbox(slide, 10.05, 1.15, 2.16, 0.55, "8s", size=42, color=RGBColor(255, 255, 255), bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 10.05, 1.82, 2.16, 0.4, "快速响应目标", size=13, color=RGBColor(255, 255, 255), align=PP_ALIGN.CENTER)
    for i, item in enumerate(["五模型并行", "置信度投票", "真实法规库", "多场景扩展"]):
        add_rect(slide, 10.02, 2.65 + i * 0.65, 2.18, 0.42, fill=RGBColor(255, 255, 255), radius=True)
        add_textbox(slide, 10.14, 2.74 + i * 0.65, 1.94, 0.2, item, size=11, color=FUDAN_RED, bold=True, align=PP_ALIGN.CENTER)


def content_title(slide, title: str, section: str | None = None) -> None:
    add_rect(slide, 0, 0, 13.33, 7.5, fill=LIGHT_BG)
    add_rect(slide, 0, 0, 13.33, 0.18, fill=FUDAN_RED)
    if section:
        add_badge(slide, 0.55, 0.42, section, fill=PALE_RED, color=FUDAN_RED, w=1.55)
    add_textbox(slide, 0.55, 0.78, 10.8, 0.45, title, size=24, color=FUDAN_RED, bold=True)
    add_rect(slide, 0.55, 1.34, 1.1, 0.04, fill=FUDAN_GOLD)


def add_bullets(slide, x, y, w, h, items, size=15, color=FUDAN_DARK, bullet_color=FUDAN_RED):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(6)
        run = p.add_run()
        run.text = f"• {item}"
        run.font.name = "PingFang SC"
        run.font.size = Pt(size)
        run.font.color.rgb = color
    return box


def add_card(slide, x, y, w, h, title, body, accent=FUDAN_RED, body_size=12):
    add_rect(slide, x, y, w, h, fill=PANEL_BG, line=RGBColor(230, 224, 224), radius=True)
    add_rect(slide, x, y, 0.08, h, fill=accent)
    add_textbox(slide, x + 0.2, y + 0.16, w - 0.35, 0.28, title, size=14, color=accent, bold=True)
    add_textbox(slide, x + 0.2, y + 0.55, w - 0.35, h - 0.68, body, size=body_size, color=FUDAN_DARK, line_spacing=1.05)


def add_metric(slide, x, y, value, label, color=FUDAN_RED):
    add_rect(slide, x, y, 2.35, 1.16, fill=PANEL_BG, line=RGBColor(232, 226, 226), radius=True)
    add_textbox(slide, x + 0.1, y + 0.16, 2.1, 0.42, value, size=26, color=color, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, x + 0.1, y + 0.68, 2.1, 0.24, label, size=10, color=FUDAN_GRAY, align=PP_ALIGN.CENTER)


def add_arrow(slide, x1, y1, x2, y2, color=FUDAN_RED):
    line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(1.5)
    return line


def add_process(slide, y, steps):
    box_w = 1.72
    gap = 0.28
    x = 0.62
    for i, (label, sub) in enumerate(steps):
        add_rect(slide, x, y, box_w, 1.02, fill=PANEL_BG, line=RGBColor(226, 220, 220), radius=True)
        add_textbox(slide, x + 0.12, y + 0.18, box_w - 0.24, 0.22, label, size=11, color=FUDAN_RED, bold=True, align=PP_ALIGN.CENTER)
        add_textbox(slide, x + 0.13, y + 0.50, box_w - 0.26, 0.32, sub, size=8, color=FUDAN_GRAY, align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            add_arrow(slide, x + box_w + 0.03, y + 0.51, x + box_w + gap - 0.03, y + 0.51, color=FUDAN_GOLD)
        x += box_w + gap


def add_table(slide, x, y, col_widths, rows, header=True, font_size=10):
    table = slide.shapes.add_table(len(rows), len(col_widths), Inches(x), Inches(y), Inches(sum(col_widths)), Inches(0.35 * len(rows))).table
    for c, width in enumerate(col_widths):
        table.columns[c].width = Inches(width)
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(text)
            cell.margin_left = Inches(0.04)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            fill = FUDAN_RED if header and r == 0 else (RGBColor(253, 250, 245) if r % 2 else RGBColor(255, 255, 255))
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER if c != 1 else PP_ALIGN.LEFT
                for run in p.runs:
                    run.font.name = "PingFang SC"
                    run.font.size = Pt(font_size)
                    run.font.bold = header and r == 0
                    run.font.color.rgb = RGBColor(255, 255, 255) if header and r == 0 else FUDAN_DARK
    return table


def add_image_fit(slide, path: Path, x, y, w, h):
    if not path.exists():
        add_rect(slide, x, y, w, h, fill=RGBColor(230, 230, 230), line=RGBColor(210, 210, 210), radius=True)
        add_textbox(slide, x + 0.2, y + h / 2 - 0.15, w - 0.4, 0.3, f"图片缺失：{path.name}", size=10, color=FUDAN_GRAY, align=PP_ALIGN.CENTER)
        return
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
    if pic.height > Inches(h):
        pic.height = Inches(h)
    if pic.width > Inches(w):
        pic.width = Inches(w)
    pic.left = Inches(x + (w - pic.width / 914400) / 2)
    pic.top = Inches(y + (h - pic.height / 914400) / 2)


def create_deck() -> Path:
    prs = Presentation(str(TEMPLATE))
    delete_all_slides(prs)
    total = 15
    today = datetime.now().strftime("%Y/%m/%d")
    blank = prs.slide_layouts[4]

    slides = []
    for _ in range(total):
        slides.append(prs.slides.add_slide(blank))

    title_slide(slides[0], "组会汇报", TITLE, SUBTITLE, f"汇报日期：{today}")

    content_title(slides[1], "汇报目录", "CONTENTS")
    toc = [
        ("01", "任务背景与目标", "老师要求：非网页版、API脚本、8秒内输出风险与条例"),
        ("02", "系统设计与实现", "图片预处理、Prompt协议、五模型并行、投票融合"),
        ("03", "法规条款与结果可信度", "citation_key法规库、证据门控、实名投票"),
        ("04", "真实测试与对比", "临时用电图片、同组报告对比、耗时拆解"),
        ("05", "问题复盘与下一步", "超时、重复风险、模型波动与工程优化"),
    ]
    for i, (num, head, body) in enumerate(toc):
        y = 1.65 + i * 0.9
        add_textbox(slides[1], 0.85, y, 0.7, 0.35, num, size=17, color=FUDAN_RED, bold=True, align=PP_ALIGN.CENTER)
        add_textbox(slides[1], 1.72, y - 0.02, 3.2, 0.28, head, size=15, color=FUDAN_DARK, bold=True)
        add_textbox(slides[1], 5.0, y - 0.01, 6.7, 0.28, body, size=12, color=FUDAN_GRAY)
        add_rect(slides[1], 1.63, y + 0.45, 10.2, 0.01, fill=RGBColor(225, 218, 218))
    add_footer(slides[1], 2, total)

    content_title(slides[2], "项目背景：工地安全检测从“慢复核”走向“快初筛”", "01")
    add_card(slides[2], 0.7, 1.65, 3.7, 4.5, "痛点", "工地现场复杂：遮挡、反光、模糊、空间关系难判断；人工逐图检查慢，且标准条款检索成本高。", accent=FUDAN_RED, body_size=14)
    add_card(slides[2], 4.8, 1.65, 3.7, 4.5, "目标", "输入单张施工现场图片，在8秒内完成安全风险快速识别，给出违规点、图像证据、真实可查安全条例。", accent=FUDAN_GOLD, body_size=14)
    add_card(slides[2], 8.9, 1.65, 3.7, 4.5, "边界", "定位为快速初筛工具，不替代现场测量、设备功能试验、资料审查和专业人员最终验收。", accent=BLUE, body_size=14)
    add_footer(slides[2], 3, total)

    content_title(slides[3], "任务演进：从单模型 Prompt 到多模型投票系统", "01")
    rows = [
        ("阶段", "主要问题", "对应改进"),
        ("1. 单次调用", "DeepSeek Flash可快速返回，但低思考模式误判较多", "压缩输出、关闭思考、限制证据边界"),
        ("2. Prompt增强", "空间位置、光影、箱门/插座盖等容易错判", "强调可见性、自检、部件拆分识别"),
        ("3. 双模型复检", "单模型不稳定，存在偶发误检", "引入Qwen视觉二检"),
        ("4. 五模型并行", "主辅模型关系不清，输出可信度难解释", "DeepSeek/Qwen/GLM/Kimi/豆包平权投票"),
        ("5. 法规库", "模型可能编造条款编号", "citation_key白名单映射真实法规"),
    ]
    add_table(slides[3], 0.65, 1.55, [1.55, 4.0, 6.1], rows, font_size=9)
    add_footer(slides[3], 4, total)

    content_title(slides[4], "系统总体架构：代码编排，大模型负责视觉安全判断", "02")
    steps = [
        ("输入图片", "场景参数"),
        ("图片预处理", "压缩/质量分析"),
        ("统一Prompt", "证据边界/条款键"),
        ("五模型并行", "同图独立检测"),
        ("JSON规范化", "状态/证据字段"),
        ("投票融合", "实名投票/置信度"),
        ("法规补全", "真实条款输出"),
    ]
    add_process(slides[4], 1.75, steps)
    add_bullets(slides[4], 0.85, 3.4, 5.65, 1.7, [
        "核心原则：不让代码用有限规则替代大模型识图判断。",
        "代码负责工程约束：超时、结构化、去重、评分、保存报告。",
        "模型负责主要判断：识别对象、发现风险、给出图片证据。"
    ], size=14)
    add_card(slides[4], 7.1, 3.28, 5.25, 1.95, "当前主流程", "load .env → prepare_image_data_url → ThreadPoolExecutor并行请求 → extract_json_object → strict_issue_quality_gate → merge_model_issues → model_judged_report", accent=FUDAN_RED, body_size=11)
    add_footer(slides[4], 5, total)

    content_title(slides[5], "模型方案：五个国产大模型平权并行", "02")
    rows = [
        ("模型", "定位", "优势", "约束"),
        ("DeepSeek Flash", "快速视觉检测", "响应快，清晰目标识别稳定", "低思考模式细节偶有误判"),
        ("Qwen3.8-Flash", "独立视觉检测", "速度稳定，文本结构较好", "对局部细节偶尔保守"),
        ("GLM-5.3-Flash", "异构视觉检测", "提供不同厂商判断", "网络波动时可能超时"),
        ("Kimi K2.6", "独立视觉检测", "能补充部件识别", "通常是最慢一路"),
        ("豆包 Lite", "低延迟视觉检测", "较快，适合投票扩展", "依赖火山方舟接入点"),
    ]
    add_table(slides[5], 0.58, 1.45, [1.7, 1.9, 3.8, 4.0], rows, font_size=8)
    add_metric(slides[5], 0.9, 5.95, "5路", "并行检测")
    add_metric(slides[5], 3.55, 5.95, "同一Prompt", "减少口径差异", color=BLUE)
    add_metric(slides[5], 6.2, 5.95, "实名投票", "每条风险可追溯", color=GREEN)
    add_metric(slides[5], 8.85, 5.95, "8秒", "总时间预算", color=ORANGE)
    add_footer(slides[5], 6, total)

    content_title(slides[6], "Prompt设计：高精度检查协议而不是简单关键词限制", "02")
    add_card(slides[6], 0.65, 1.45, 3.9, 4.9, "证据边界", "只依据清晰可见的静态证据；未入镜、遮挡、反光、模糊、过小或需要测量/动态验证的内容，不得判为明确违规。", accent=FUDAN_RED, body_size=13)
    add_card(slides[6], 4.8, 1.45, 3.9, 4.9, "易错对象拆分", "配电箱主箱门、插座盖、插头、端子、电缆接口分别判断；不把一个部件状态复制给另一个部件。", accent=FUDAN_GOLD, body_size=13)
    add_card(slides[6], 8.95, 1.45, 3.65, 4.9, "结构化输出", "每个风险必须返回 status、risk_key、evidence、confidence、target_visibility、evidence_level、citation_key。", accent=BLUE, body_size=12)
    add_footer(slides[6], 7, total)

    content_title(slides[7], "置信度与投票融合：把“模型说了什么”变成可解释证据", "03")
    add_card(slides[7], 0.68, 1.45, 5.65, 2.1, "单模型得分", "0.45 × 状态分 + 0.35 × 模型自报置信度 + 0.20 × 证据质量分\n\n状态分区分 CLEAR / SUSPECTED / NOT_ASSESSABLE；证据质量由 direct / partial / ambiguous 等字段决定。", accent=FUDAN_RED, body_size=12)
    add_card(slides[7], 6.75, 1.45, 5.65, 2.1, "融合置信度", "支持模型平均分 + 投票加成\n\n达到多数支持时作为主要风险；1/5 或 2/5 会被标注为“疑似、需复核”，避免单模型误判直接变成确定结论。", accent=GREEN, body_size=12)
    rows = [
        ("输出字段", "含义"),
        ("支持", "该模型独立识别出同一根本风险"),
        ("未支持", "模型完成检测，但没有提出该风险"),
        ("未返回", "调用失败、超时或JSON解析失败"),
        ("实名投票", "逐项显示每个模型的投票状态和单模型得分"),
    ]
    add_table(slides[7], 1.2, 4.15, [2.5, 7.9], rows, font_size=10)
    add_footer(slides[7], 8, total)

    content_title(slides[8], "真实法规输出：解决大模型“编造条款”的风险", "03")
    add_bullets(slides[8], 0.85, 1.48, 5.55, 2.0, [
        "模型不再直接生成最终法规文本，只能选择 citation_key。",
        "Python从 verified_regulations.py 白名单读取标准名称、条款号、原文和查询链接。",
        "未知键统一降级为“暂未匹配到已核验的具体条文”。"
    ], size=14)
    code = (
        "TEMP_ELEC_CABLE_LAYOUT → JGJ/T 46-2024 第6.2.3条\n"
        "TEMP_ELEC_WIRING_ENTRY → JGJ/T 46-2024 第4.1.15条\n"
        "TEMP_ELEC_BOX_CLEAN → JGJ/T 46-2024 第4.3.8条\n"
        "TEMP_ELEC_BOX_SPACE → JGJ/T 46-2024 第4.1.5条"
    )
    add_card(slides[8], 6.85, 1.35, 5.65, 2.55, "条款键示例", code, accent=FUDAN_RED, body_size=12)
    add_card(slides[8], 0.85, 4.35, 11.65, 1.35, "效果", "报告中的“违反的具体安全条例”来自本地核验库，而不是模型自由编写。这样可以在组会、论文或答辩中明确说明条款来源可查，降低法规幻觉风险。", accent=FUDAN_GOLD, body_size=13)
    add_footer(slides[8], 9, total)

    content_title(slides[9], "代码结构：从命令行到Markdown报告的闭环", "03")
    rows = [
        ("模块/文件", "职责"),
        ("safety_check.py", "参数解析、图片处理、Prompt构造、API并行调用、投票融合、报告保存"),
        ("scene_rules.py", "多场景目录、场景别名、检查字段和场景提示语"),
        ("verified_regulations.py", "已核验法规条款库，按 citation_key 输出真实条文"),
        (".env / .env.example", "模型API、超时、图片压缩、token预算等配置"),
        ("outputs/", "保存每次检测生成的Markdown报告，便于复盘与对比"),
    ]
    add_table(slides[9], 0.65, 1.45, [2.65, 8.75], rows, font_size=9)
    add_card(slides[9], 1.0, 5.58, 10.9, 0.85, "运行方式", '.venv/bin/python safety_check.py "/图片路径" --scene "施工现场临时用电" --show-timing', accent=BLUE, body_size=12)
    add_footer(slides[9], 10, total)

    content_title(slides[10], "多场景适配：从临时用电扩展到复杂工地安全", "03")
    scenes = [
        ("综合安全", "PPE / 通道 / 消防 / 材料堆放"),
        ("临时用电", "配电箱 / 电缆 / 接地 / 漏保"),
        ("高处作业", "临边洞口 / 安全带 / 梯子"),
        ("基坑工程", "支护 / 临边 / 排水 / 堆载"),
        ("起重吊装", "吊物 / 吊具 / 警戒 / 站位"),
        ("塔吊/升降机", "附着 / 层门 / 限位 / 通道"),
        ("脚手架", "连墙件 / 剪刀撑 / 脚手板"),
        ("扩展机制", "新增场景配置与法规键即可迭代"),
    ]
    for i, (head, body) in enumerate(scenes):
        x = 0.68 + (i % 4) * 3.13
        y = 1.45 + (i // 4) * 2.05
        add_card(slides[10], x, y, 2.75, 1.62, head, body, accent=[FUDAN_RED, FUDAN_GOLD, BLUE, GREEN][i % 4], body_size=10)
    add_footer(slides[10], 11, total)

    content_title(slides[11], "真实测试一：打开配电箱内部图，与同组完整版报告对比", "04")
    add_image_fit(slides[11], IMG1, 0.65, 1.45, 5.25, 3.95)
    rows = [
        ("风险", "本系统", "同组报告"),
        ("进出线口无护套/固定线卡", "4/5支持", "确认违规"),
        ("箱内杂物", "4/5支持", "确认违规"),
        ("线束无外套绝缘管", "1/5支持", "确认违规"),
        ("用途标识/系统图缺失", "未稳定检出", "确认违规"),
    ]
    add_table(slides[11], 6.18, 1.55, [2.8, 1.55, 2.25], rows, font_size=8)
    add_card(slides[11], 6.18, 4.55, 6.65, 1.1, "结论", "快速系统在最明显风险上与完整版报告一致；对需要近景、门板或铭牌确认的细项更保守，标记为疑似或未输出。", accent=FUDAN_RED, body_size=12)
    add_footer(slides[11], 12, total)

    content_title(slides[12], "真实测试二：室外关闭配电箱图，电缆拖地稳定检出", "04")
    add_image_fit(slides[12], IMG2, 0.65, 1.45, 5.4, 3.95)
    rows = [
        ("风险项", "投票", "条例"),
        ("电缆沿地面明设拖地", "5/5", "JGJ/T 46-2024 第6.2.3条"),
        ("操作空间被占用", "3/5", "JGJ/T 46-2024 第4.1.5条"),
        ("进出线口不规范", "2/5", "JGJ/T 46-2024 第4.1.15条"),
    ]
    add_table(slides[12], 6.35, 1.55, [3.0, 1.1, 2.65], rows, font_size=8)
    add_metric(slides[12], 6.55, 4.55, "7.69s", "五模型总耗时")
    add_metric(slides[12], 9.25, 4.55, "5/5", "拖地电缆一致支持", color=GREEN)
    add_footer(slides[12], 13, total)

    content_title(slides[13], "性能与现象复盘：速度达标，但输出稳定性仍需工程加固", "05")
    rows = [
        ("环节", "典型耗时", "说明"),
        ("图片预处理", "0.4-0.9s", "PNG/JPG压缩、质量分析、Base64转换"),
        ("DeepSeek", "1.4-1.8s", "最快一路，负责独立视觉判断"),
        ("Qwen/豆包", "3.0-3.8s", "较稳定，适合作为快速投票成员"),
        ("GLM", "3.5-6.0s", "偶发超时，受服务状态影响"),
        ("Kimi", "6.0-7.1s", "常为最慢一路，决定总等待时间"),
        ("融合/报告", "<0.01s", "解析、打分、投票、格式化几乎不耗时"),
    ]
    add_table(slides[13], 0.65, 1.45, [2.1, 1.55, 7.8], rows, font_size=9)
    add_card(
        slides[13],
        1.0,
        5.85,
        10.9,
        0.7,
        "现象解释",
        "同一图片每次结果不完全一致，主要来自在线模型排队、超时、生成随机性和同类风险合并粒度差异。",
        accent=ORANGE,
        body_size=11,
    )
    add_footer(slides[13], 14, total)

    content_title(slides[14], "工作成果与下一步计划", "05")
    add_card(slides[14], 0.75, 1.35, 5.55, 4.7, "已完成成果", "• 完成API脚本，支持VS Code命令行运行\n• 支持五模型平权并行图片检测\n• 设计高精度Prompt与证据门控\n• 实现实名投票、置信度评分和耗时统计\n• 建立JGJ/T 46-2024真实条款库\n• 生成README、流程文档和可保存Markdown报告", accent=FUDAN_RED, body_size=13)
    add_card(slides[14], 6.85, 1.35, 5.55, 4.7, "下一步优化", "• 强化同类风险去重，尤其同一citation_key重复项\n• 对超时模型做降级策略与缓存策略\n• 扩展高处、吊装、脚手架、基坑等法规库\n• 建立小规模标注测试集，量化准确率/召回率\n• 增加近景补拍建议和人工复核入口", accent=GREEN, body_size=13)
    add_textbox(slides[14], 0.9, 6.55, 11.6, 0.35, "目标：保留大模型开放识图能力，同时用工程化流程把输出变得更快、更稳、更可解释。", size=15, color=FUDAN_RED, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slides[14], 15, total)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUT_DIR / f"{datetime.now().strftime('%Y%m%d-%H%M%S')}-复杂环境安全巡检组会汇报.pptx"
    prs.save(out)
    return out


if __name__ == "__main__":
    path = create_deck()
    print(path)
